from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain.text_splitter import RecursiveCharacterTextSplitter
import docx
from typing import List

def load_document(file_path: str):
    """根据文件后缀解析文档"""
    if file_path.endswith(".pdf"):
        loader = PyPDFLoader(file_path)
        return loader.load()
    elif file_path.endswith(".docx"):
        document = docx.Document(file_path)
        text = "\n".join([p.text for p in document.paragraphs if p.text.strip()])
        # 包装成 LangChain 文档格式
        from langchain.schema import Document
        return [Document(page_content=text, metadata={"source": file_path})]
    elif file_path.endswith(".txt"):
        loader = TextLoader(file_path, encoding="utf-8")
        return loader.load()
    else:
        raise ValueError("不支持的文件格式，仅支持 PDF、Word(docx)、TXT")

def split_documents(documents, chunk_size=500, chunk_overlap=50):
    """将文档切分成小块"""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", "。", "！", "？", " ", ""]
    )
    return splitter.split_documents(documents)